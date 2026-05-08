from __future__ import annotations

import io
import os
from typing import Any

import csv

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover
    PdfReader = None

try:
    import docx
except Exception:  # pragma: no cover
    docx = None


def _safe_decode(data: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_text_from_upload(uploaded_file: Any) -> tuple[str, str]:
    """Extract readable text and a preview note from TXT, PDF, DOCX, or CSV upload."""
    if uploaded_file is None:
        return "", "No file uploaded."
    name = uploaded_file.name
    ext = os.path.splitext(name)[1].lower()
    data = uploaded_file.getvalue()

    if ext == ".txt":
        text = _safe_decode(data)
        return text, f"TXT file extracted successfully. Characters: {len(text):,}"

    if ext == ".pdf":
        if PdfReader is None:
            return "", "PDF support is unavailable. Install pypdf."
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = []
            for i, page in enumerate(reader.pages):
                try:
                    pages.append(page.extract_text() or "")
                except Exception:
                    pages.append("")
            text = "\n\n".join(pages).strip()
            return text, f"PDF extracted successfully. Pages: {len(reader.pages)}. Characters: {len(text):,}"
        except Exception as exc:
            return "", f"Could not read PDF: {exc}"

    if ext == ".docx":
        if docx is None:
            return "", "DOCX support is unavailable. Install python-docx."
        try:
            document = docx.Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts).strip()
            return text, f"DOCX extracted successfully. Characters: {len(text):,}"
        except Exception as exc:
            return "", f"Could not read DOCX: {exc}"

    if ext == ".csv":
        try:
            decoded = _safe_decode(data)
            reader = list(csv.reader(io.StringIO(decoded)))
            if not reader:
                return "", "CSV file is empty."
            rows_as_text = [" | ".join(row) for row in reader]
            preview = "\n".join(rows_as_text[:20])
            text = "\n".join(rows_as_text)
            cols = len(reader[0]) if reader and reader[0] else 0
            return text, f"CSV extracted successfully. Rows: {max(len(reader)-1, 0):,}. Columns: {cols:,}. Preview rows included below.\n\n{preview}"
        except Exception as exc:
            return "", f"Could not read CSV: {exc}"

    return "", "Unsupported file format. Please upload TXT, PDF, DOCX, or CSV."
